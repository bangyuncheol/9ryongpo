"""배마중_발표.pdf 의 쪽을 이미지로 떠서 PPT 로 만든다. 영상 장(2·5)은 영상을 바로 재생할 수 있게 넣는다.

쓰는 법(윈도우·맥 공통):
  uv run --with pymupdf,python-pptx python make_pptx.py [소개영상.mp4] [취재영상.mp4]
- 기본 영상: brag-output/final/배마중_소개영상_최종.mp4(완성본, H.264 1080p), brag-output/배마중_문제인식영상.mp4. HEVC 파일을 주면 H.264 사본으로 바꿔 넣는다
- 쪽 번호와 영상 장 번호는 배마중_발표.html 의 장 순서를 따른다(VIDEO_PAGES).
- 각 장 노트에는 배마중_발표_노트.md 의 발표 메모와 출처가 들어간다.
- 결과 03_발표/배마중_발표.pptx 는 영상이 들어 있어 100MB 를 넘으므로 git 에 올리지 않는다(.gitignore).
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PDF = HERE / "배마중_발표.pdf"
NOTES = HERE / "배마중_발표_노트.md"
OUT = HERE / "배마중_발표.pptx"
VIDEO_PAGES = {2: ROOT / "brag-output" / "final" / "배마중_소개영상_최종.mp4", 5: ROOT / "brag-output" / "배마중_문제인식영상.mp4"}
W, H = Emu(12192000), Emu(6858000)  # 13.333 × 7.5 in (16:9)


def codec(path: Path) -> str:
    out = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=codec_name", "-of", "csv=p=0", str(path)],
                         capture_output=True, text=True, check=True).stdout.strip()
    return out


def h264_copy(src: Path, tmp: Path) -> Path:
    """파워포인트(윈도우)는 HEVC 를 기본으로 못 틀어서 H.264 로 바꾼 사본을 쓴다. 이미 H.264 면 그대로."""
    if codec(src) == "h264":
        return src
    dst = tmp / (src.stem + "_h264.mp4")
    print(f"{src.name}: {codec(src)} → H.264 변환 중")
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(src), "-c:v", "libx264", "-preset", "fast", "-crf", "20", "-pix_fmt", "yuv420p",
                    "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart", str(dst)], check=True)
    return dst


def load_notes() -> dict[int, str]:
    """노트 md → {쪽: '발표 메모 … \n\n출처: …'}"""
    notes: dict[int, str] = {}
    if not NOTES.exists():
        return notes
    text = NOTES.read_text(encoding="utf-8")
    for m in re.finditer(r"^## (\d+)\. .*?\n(.*?)(?=^## \d+\. |\Z)", text, re.S | re.M):
        body = m.group(2)
        memo = re.search(r"\*\*발표 메모\*\*\n\n(.*?)\n\n\*\*출처\*\*", body, re.S)
        src = re.search(r"\*\*출처\*\*\n\n(.*?)\s*$", body, re.S)
        parts = [memo.group(1).strip() if memo else ""]
        if src and src.group(1).strip() != "(없음)":
            parts.append("출처: " + src.group(1).strip())
        notes[int(m.group(1))] = "\n\n".join(p for p in parts if p and p != "(없음)")
    return notes


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    videos = dict(VIDEO_PAGES)
    if len(sys.argv) > 1:
        videos[2] = Path(sys.argv[1])
    if len(sys.argv) > 2:
        videos[5] = Path(sys.argv[2])
    notes = load_notes()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        doc = pymupdf.open(PDF)
        for i, page in enumerate(doc, start=1):
            png = tmp / f"page-{i:02d}.png"
            page.get_pixmap(dpi=144).save(png)
            slide = prs.slides.add_slide(blank)
            if i in videos and videos[i].exists():
                mp4 = h264_copy(videos[i], tmp)
                slide.shapes.add_movie(str(mp4), 0, 0, W, H, poster_frame_image=str(png), mime_type="video/mp4")
                print(f"{i:2d}장: 영상 {videos[i].name}")
            else:
                if i in videos:
                    print(f"{i:2d}장: 영상 파일 없음 → 이미지만 ({videos[i]})")
                slide.shapes.add_picture(str(png), 0, 0, W, H)
            if notes.get(i):
                slide.notes_slide.notes_text_frame.text = notes[i]
        out = OUT
        try:
            prs.save(out)
        except PermissionError:
            # 파워포인트에서 열어 둔 파일은 덮어쓸 수 없다. 새 이름으로 저장하고 알린다
            out = OUT.with_name(OUT.stem + "_new.pptx")
            prs.save(out)
            print(f"{OUT.name} 이 열려 있어 덮어쓰지 못했습니다. 닫은 뒤 {out.name} 을 {OUT.name} 으로 바꿔 쓰세요.")
    print(f"{out.name}: {len(prs.slides)}장, {out.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
